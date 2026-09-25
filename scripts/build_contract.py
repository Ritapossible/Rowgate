"""Generate contract/api-contract.xlsx, the signed partner contract Rowgate checks against.

The workbook is deliberately untidy, like the real thing: every sheet has a different
layout, headers are merged, one status code is a merged block, rules reference each
other and the shared Errors sheet, and some rows are PLANNED or DEPRECATED and must not
be enforced.

Run:  python scripts/build_contract.py
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).resolve().parent.parent / "contract" / "api-contract.xlsx"

TITLE = Font(bold=True, size=14, color="FFFFFF")
TITLE_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
GROUP_FILL = PatternFill("solid", fgColor="B4C6E7")
GREY = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def title(ws, text, span):
    ws.merge_cells(f"A1:{span}1")
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = CENTER
    ws.row_dimensions[1].height = 26


def header(ws, row, labels):
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.font = HEAD
        c.fill = HEAD_FILL
        c.border = BOX
        c.alignment = CENTER


def rows(ws, first_row, data, status_col):
    for r, values in enumerate(data, start=first_row):
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.border = BOX
            c.alignment = WRAP
        if ws.cell(row=r, column=status_col).value in ("PLANNED", "DEPRECATED"):
            for col in range(1, len(values) + 1):
                ws.cell(row=r, column=col).fill = GREY


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def cover(wb):
    ws = wb.active
    ws.title = "Cover"
    title(ws, "Kora POS ⇄ Kora Partner API — Integration Contract", "F")
    info = [
        ("Partner", "Kora Retail Ltd. (POS and mobile checkout)"),
        ("Provider", "Kora Partner API team"),
        ("Contract version", "v3.2"),
        ("Signed", "2026-08-14"),
        ("Scope", "Orders, Billing, Auth. Webhooks are covered by a separate contract (KWH-1)."),
        ("Change rule", "This workbook is the signed source of truth. Any behaviour change to an "
                        "ACTIVE row needs a Changelog entry and partner sign-off before release."),
        ("How to read", "Rows marked PLANNED or DEPRECATED are informational and not enforced. "
                        "\"↳ ORD-011\" means the value is inherited from that rule."),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = HEAD
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
        ws.cell(row=i, column=2, value=v).alignment = WRAP
        ws.row_dimensions[i].height = 32
    ws.merge_cells("A11:F11")
    ws["A11"] = "Sign-off"
    ws["A11"].font = HEAD
    ws["A11"].fill = GROUP_FILL
    header(ws, 12, ["Party", "Name", "Role", "Date", "Signature", ""])
    rows(ws, 13, [
        ["Kora Retail", "A. Okafor", "Head of Payments", "2026-08-14", "signed", ""],
        ["Kora Partner API", "T. Bello", "Engineering Manager", "2026-08-14", "signed", ""],
    ], status_col=6)
    widths(ws, {"A": 20, "B": 26, "C": 22, "D": 14, "E": 14, "F": 20})


def orders(wb):
    ws = wb.create_sheet("Orders")
    title(ws, "Orders — POST /orders, GET /orders/{id}", "I")
    # Row 2: merged group headers
    for span, label in (("A2:C2", "Operation"), ("D2:G2", "Response contract"), ("H2:I2", "Governance")):
        ws.merge_cells(span)
        c = ws[span.split(":")[0]]
        c.value = label
        c.font = HEAD
        c.fill = GROUP_FILL
        c.alignment = CENTER
    header(ws, 3, ["Rule", "Operation", "HTTP status", "Field", "Type", "Required",
                   "Scenario / error ref", "Contract status", "Notes"])
    data = [
        # row 4
        ["ORD-001", "POST /orders", "↳ ORD-011", "order_id", "string (ord_*)", "Y", "valid order", "ACTIVE", ""],
        ["ORD-002", "POST /orders", "↳ ORD-011", "status", "enum: pending | confirmed", "Y", "valid order", "ACTIVE", ""],
        ["ORD-003", "POST /orders", "↳ ORD-011", "total_cents", "integer", "Y", "valid order", "ACTIVE",
         "Unit price × quantity, minor units."],
        ["ORD-004", "POST /orders", "↳ ORD-011", "currency", "string (ISO 4217)", "Y", "valid order", "ACTIVE", ""],
        ["ORD-005", "POST /orders", "↳ ORD-011", "created_at", "string (RFC 3339, UTC)", "Y", "valid order", "ACTIVE", ""],
        # row 9
        ["ORD-006", "POST /orders", 422, "error", "string", "Y", "quantity ≤ 0 or > 100 → ERR-002", "ACTIVE",
         "Body per Errors sheet."],
        ["ORD-007", "POST /orders", 409, "error", "string", "Y", "duplicate Idempotency-Key → ERR-003", "PLANNED",
         "Targeted for v3.3. Do not test yet."],
        # row 11: request_id
        ["ORD-008", "POST /orders", "↳ ORD-011", "request_id", "string", "Y",
         "valid order; echoes X-Request-ID if sent", "ACTIVE",
         "POS reconciles orders by request_id. Name on the wire must be exactly request_id."],
        ["ORD-009", "GET /orders/{id}", 200, "order_id", "string", "Y", "known id", "ACTIVE", ""],
        ["ORD-010", "GET /orders/{id}", 200, "status", "enum: pending | confirmed", "Y", "known id", "ACTIVE", ""],
        # row 14: the status code rule for create
        ["ORD-011", "POST /orders", 201, "—", "—", "—", "valid order: synchronous create", "ACTIVE",
         "POS treats any other 2xx as a failed checkout and retries (see INC-2291)."],
        ["ORD-012", "GET /orders/{id}", 404, "error", "string", "Y", "unknown id → ERR-005", "ACTIVE", ""],
        ["ORD-013", "DELETE /orders/{id}", 204, "—", "—", "—", "cancel", "DEPRECATED",
         "Removed in v3.0. Cancellation moves to POST /orders/{id}/cancel."],
        ["ORD-014", "POST /orders/{id}/cancel", 200, "status", "enum: cancelled", "Y", "cancel", "PLANNED", ""],
        ["ORD-015", "GET /orders/{id}", 200, "total_cents", "integer", "Y", "known id", "ACTIVE", ""],
        ["ORD-016", "GET /orders/{id}", 200, "currency", "string (ISO 4217)", "Y", "known id", "ACTIVE", ""],
    ]
    rows(ws, 4, data, status_col=8)
    ws["I14"].comment = Comment("Agreed after INC-2291 (duplicate charges, July 2026).", "A. Okafor")
    widths(ws, {"A": 10, "B": 24, "C": 12, "D": 14, "E": 24, "F": 9, "G": 34, "H": 15, "I": 48})
    ws.freeze_panes = "A4"
    return ws


def billing(wb):
    ws = wb.create_sheet("Billing")
    title(ws, "Billing — GET /invoices/{id}", "I")
    header(ws, 2, ["Rule", "Endpoint", "Scenario", "HTTP status", "Field", "Type", "Required",
                   "Contract status", "Notes"])
    ep = "GET /invoices/{id}"
    data = [
        # rows 3..9 share one merged HTTP status cell (D3:D9)
        ["BIL-001", ep, "known invoice", 200, "invoice_id", "string (inv_*)", "Y", "ACTIVE", ""],
        ["BIL-002", ep, "known invoice", None, "order_id", "string", "Y", "ACTIVE", ""],
        ["BIL-003", ep, "known invoice", None, "amount_cents", "integer", "Y", "ACTIVE", "Equals order total_cents."],
        ["BIL-004", ep, "known invoice", None, "issued_at", "string (RFC 3339)", "Y", "ACTIVE", ""],
        ["BIL-005", ep, "known invoice", None, "status", "enum: open | paid | void", "Y", "ACTIVE", ""],
        ["BIL-006", ep, "known invoice", None, "due_date", "string (YYYY-MM-DD)", "Y", "ACTIVE", "issued_at + 30 days."],
        # row 9: currency
        ["BIL-007", ep, "known invoice", None, "currency", "string (ISO 4217)", "Y", "ACTIVE",
         "Mandatory for multi-currency settlement. Partner ledger rejects invoices without it."],
        ["BIL-008", ep, "unknown id", 404, "error", "string", "Y", "ACTIVE", "→ ERR-005"],
        ["BIL-009", ep, "known invoice", 200, "tax_cents", "integer", "N", "ACTIVE", "Optional. VAT 7.5% when present."],
        ["BIL-010", "GET /invoices?order_id=", "list by order", 200, "items[]", "array", "Y", "PLANNED", "v3.3"],
        ["BIL-011", "POST /invoices/{id}/void", "void", 200, "status", "enum: void", "Y", "PLANNED", "v3.4"],
    ]
    rows(ws, 3, data, status_col=8)
    ws.merge_cells("D3:D9")
    ws["D3"].alignment = CENTER
    widths(ws, {"A": 10, "B": 24, "C": 16, "D": 12, "E": 14, "F": 24, "G": 9, "H": 15, "I": 50})
    ws.freeze_panes = "A3"
    return ws


def auth(wb):
    ws = wb.create_sheet("Auth")
    title(ws, "Auth — POST /auth/token (OAuth 2.0 client credentials)", "H")
    ws["A2"] = "Errors use the shared envelope on the Errors sheet. FastAPI's default {\"detail\": …} is not accepted by POS."
    ws.merge_cells("A2:H2")
    ws["A2"].font = Font(italic=True)
    header(ws, 3, ["Rule", "Endpoint", "grant_type", "Scenario", "HTTP status", "Error ref",
                   "Contract status", "Notes"])
    data = [
        # row 4
        ["AUT-001", "POST /auth/token", "client_credentials", "valid client", 200, "—", "ACTIVE",
         "Body: access_token, token_type=Bearer, expires_in=3600."],
        # row 5: bad credentials
        ["AUT-002", "POST /auth/token", "client_credentials", "wrong client_secret", 401, "ERR-004", "ACTIVE",
         "POS shows 're-enter API key' only on 401 + invalid_grant."],
        ["AUT-003", "POST /auth/token", "password", "unsupported grant", 400, "ERR-006", "ACTIVE", ""],
        ["AUT-004", "POST /auth/token", "client_credentials", "missing client_id", 422, "ERR-002", "ACTIVE", ""],
        ["AUT-005", "POST /auth/refresh", "refresh_token", "refresh", 200, "—", "PLANNED", "v3.4"],
        ["AUT-006", "POST /auth/token", "client_credentials", "valid client", 200, "—", "ACTIVE",
         "token_type must be exactly \"Bearer\"."],
    ]
    rows(ws, 4, data, status_col=7)
    widths(ws, {"A": 10, "B": 20, "C": 20, "D": 22, "E": 12, "F": 10, "G": 15, "H": 50})
    ws.freeze_panes = "A4"
    return ws


def errors(wb):
    ws = wb.create_sheet("Errors")
    title(ws, "Errors — shared error envelope", "F")
    header(ws, 2, ["Code", "HTTP status", "error", "Body (exact shape)", "Used by", "Notes"])
    body = '{{"error": "{e}", "error_description": "<text>"}}'
    data = [
        ["ERR-001", 400, "invalid_request", body.format(e="invalid_request"), "reserved", ""],
        ["ERR-002", 422, "validation_failed", body.format(e="validation_failed"), "ORD-006, AUT-004", ""],
        ["ERR-003", 409, "duplicate_request", body.format(e="duplicate_request"), "ORD-007", "PLANNED"],
        # row 6: invalid_grant
        ["ERR-004", 401, "invalid_grant",
         '{"error": "invalid_grant", "error_description": "Client authentication failed"}',
         "AUT-002", "error_description text is shown to cashiers verbatim."],
        ["ERR-005", 404, "not_found", body.format(e="not_found"), "ORD-012, BIL-008", ""],
        ["ERR-006", 400, "unsupported_grant_type", body.format(e="unsupported_grant_type"), "AUT-003", ""],
        ["ERR-007", 429, "rate_limited", body.format(e="rate_limited"), "—", "PLANNED"],
        ["ERR-008", 500, "internal_error", body.format(e="internal_error"), "all", "Never leak stack traces."],
    ]
    rows(ws, 3, data, status_col=6)
    widths(ws, {"A": 10, "B": 12, "C": 24, "D": 70, "E": 18, "F": 40})
    ws.freeze_panes = "A3"
    return ws


def changelog(wb):
    ws = wb.create_sheet("Changelog")
    title(ws, "Changelog — every change to an ACTIVE row", "F")
    header(ws, 2, ["Date", "Version", "Cell", "Change", "Decision", "Approved by"])
    rows(ws, 3, [
        ["2026-03-02", "v3.0", "Orders!A16", "DELETE /orders/{id} deprecated", "breaking, accepted", "A. Okafor"],
        ["2026-06-10", "v3.1", "Billing!E11", "tax_cents added (optional)", "additive", "T. Bello"],
        ["2026-08-14", "v3.2", "Orders!C14", "Create must return 201 (INC-2291)", "clarification", "A. Okafor"],
    ], status_col=6)
    widths(ws, {"A": 12, "B": 9, "C": 14, "D": 44, "E": 20, "F": 16})


# Cells the demo depends on. The build fails if a layout change moves them.
EXPECTED = {
    ("Orders", "A14"): "ORD-011", ("Orders", "C14"): 201,
    ("Orders", "A11"): "ORD-008", ("Orders", "D11"): "request_id",
    ("Billing", "A9"): "BIL-007", ("Billing", "E9"): "currency", ("Billing", "D3"): 200,
    ("Auth", "A5"): "AUT-002", ("Auth", "E5"): 401, ("Auth", "F5"): "ERR-004",
    ("Errors", "A6"): "ERR-004",
}


def main():
    wb = Workbook()
    cover(wb)
    orders(wb)
    billing(wb)
    auth(wb)
    errors(wb)
    changelog(wb)
    for (sheet, cell), want in EXPECTED.items():
        got = wb[sheet][cell].value
        assert got == want, f"{sheet}!{cell} is {got!r}, expected {want!r}"
    fixed = datetime(2026, 8, 14, 9, 0, 0)
    wb.properties.creator = "Kora Partner API team"
    wb.properties.created = fixed
    wb.properties.modified = fixed
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    active = sum(
        1 for name in ("Orders", "Billing", "Auth")
        for row in wb[name].iter_rows(values_only=True) if "ACTIVE" in row
    )
    total = sum(1 for name in ("Orders", "Billing", "Auth", "Errors")
                for row in wb[name].iter_rows(values_only=True)
                if isinstance(row[0], str) and row[0][:4] in ("ORD-", "BIL-", "AUT-", "ERR-"))
    print(f"wrote {OUT.relative_to(OUT.parent.parent)}: {total} rules, {active} ACTIVE")


if __name__ == "__main__":
    main()
