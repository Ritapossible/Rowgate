"""Deterministic step: render out/dossier.html from Bob's findings. No model involved.

Inputs
  out/findings.json                 written by Bob (schema: rowgate/findings.schema.json)
  out/test_results.before.json      pytest JSON report on the untouched branch (optional)
  out/test_results.after.json       pytest JSON report after the approved changes (optional)
  contract/api-contract.xlsx        the cited rows are read straight from the workbook

Usage: python scripts/render_dossier.py [--findings out/findings.json] [--out out/dossier.html]
"""

import argparse
import json
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CELL = re.compile(r"^([A-Za-z]+)!([A-Z]+[0-9]+)$")
TEST = re.compile(r"^tests/contract/test_[a-z]+_([A-Z]+[0-9]+)_[a-z0-9_]+\.py$")
HEADER_KEYS = {"Rule", "Code", "Date", "Party"}


def validate(doc: dict) -> list[str]:
    """The checks that matter for the demo. Errors are shown in the dossier, not hidden."""
    problems = []
    for key in ("branch", "base", "workbook", "findings", "skipped"):
        if key not in doc:
            problems.append(f"findings.json is missing '{key}'")
    for f in doc.get("findings", []):
        fid = f.get("id", "?")
        if not CELL.match(f.get("cell", "")):
            problems.append(f"{fid}: cell '{f.get('cell')}' is not Sheet!A1 form")
        m = TEST.match(f.get("test", ""))
        if not m:
            problems.append(f"{fid}: BREAK without a correctly named test (tests/contract/test_<sheet>_<cell>_<what>.py)")
        elif CELL.match(f.get("cell", "")) and m.group(1) != CELL.match(f["cell"]).group(2):
            problems.append(f"{fid}: test name cites {m.group(1)} but finding cites {f['cell']}")
        elif not (ROOT / f["test"]).exists():
            problems.append(f"{fid}: {f['test']} does not exist")
    return problems


def cited_row(wb, ref: str) -> dict | None:
    """Return the header row and the cited row so the dossier can show the cell in context."""
    m = CELL.match(ref)
    if not m or m.group(1) not in wb.sheetnames:
        return None
    ws = wb[m.group(1)]
    col_letter, row = coordinate_from_string(m.group(2))
    col = column_index_from_string(col_letter)
    header_row = next(
        (r for r in range(1, min(row, 10)) if ws.cell(row=r, column=1).value in HEADER_KEYS), None
    )
    width = ws.max_column

    def value(r, c):
        v = ws.cell(row=r, column=c).value
        if v is None:  # merged block: show the value from the top-left cell
            for rng in ws.merged_cells.ranges:
                if rng.min_row <= r <= rng.max_row and rng.min_col <= c <= rng.max_col:
                    v = ws.cell(row=rng.min_row, column=rng.min_col).value
                    break
        return "" if v is None else str(v)

    return {
        "sheet": m.group(1),
        "letters": [get_column_letter(c) for c in range(1, width + 1)],
        "headers": [value(header_row, c) for c in range(1, width + 1)] if header_row else [],
        "cells": [value(row, c) for c in range(1, width + 1)],
        "row": row,
        "hit": col - 1,
    }


def outcomes(path: Path) -> dict[str, str]:
    """Map test file → worst outcome in a pytest-json-report file."""
    if not path.exists():
        return {}
    report = json.loads(path.read_text())
    result: dict[str, str] = {}
    rank = {"passed": 0, "skipped": 1, "failed": 2, "error": 3}
    for t in report.get("tests", []):
        file = t["nodeid"].split("::")[0]
        outcome = t.get("outcome", "error")
        if rank.get(outcome, 3) >= rank.get(result.get(file, "passed"), 0):
            result[file] = outcome
    return result


def diff_lines(hunk: str) -> list[tuple[str, str]]:
    kinds = []
    for line in hunk.splitlines():
        kind = "add" if line.startswith("+") else "del" if line.startswith("-") else "at" if line.startswith("@@") else "ctx"
        kinds.append((kind, line))
    return kinds


def git(*args: str) -> str:
    try:
        return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()
    except (subprocess.CalledProcessError, FileNotFoundError):
        return ""


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--findings", default="out/findings.json")
    ap.add_argument("--out", default="out/dossier.html")
    args = ap.parse_args()

    doc = json.loads((ROOT / args.findings).read_text())
    problems = validate(doc)
    wb = load_workbook(ROOT / doc.get("workbook", "contract/api-contract.xlsx"))
    before = outcomes(ROOT / "out/test_results.before.json")
    after = outcomes(ROOT / "out/test_results.after.json")

    findings = []
    for f in doc.get("findings", []):
        findings.append({
            **f,
            "row": cited_row(wb, f.get("cell", "")),
            "related": [cited_row(wb, c) for c in f.get("related_cells", [])],
            "diff": diff_lines(f.get("evidence", {}).get("hunk", "")),
            "before": before.get(f.get("test", "")),
            "after": after.get(f.get("test", "")),
        })
    skipped = [{**s, "row": cited_row(wb, s.get("cell", ""))} for s in doc.get("skipped", [])]

    counts = {
        "breaks": len(findings),
        "skipped": len(skipped),
        "red_before": sum(1 for f in findings if f["before"] == "failed"),
        "green_after": sum(1 for f in findings if f["after"] == "passed"),
        "recorded": sum(1 for f in findings if f.get("decision") == "record_breaking"),
        "pending": sum(1 for f in findings if f.get("decision") == "pending"),
    }

    env = Environment(loader=FileSystemLoader(ROOT / "rowgate"), autoescape=select_autoescape(["html", "j2"]))
    html = env.get_template("dossier.html.j2").render(
        doc=doc,
        findings=findings,
        skipped=skipped,
        counts=counts,
        problems=problems,
        head=git("rev-parse", "--short", "HEAD"),
        generated=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
    )
    out = ROOT / args.out
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(html)
    print(f"wrote {args.out}: {counts['breaks']} breaks, {counts['skipped']} skipped, {len(problems)} problems")
    for p in problems:
        print(f"  ! {p}", file=sys.stderr)
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
